#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import re
import argparse
import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# 图片中的主要顺序
CATEGORY_ORDER = [
    "WEB",
    "H5",
    "全站APP",
    "代理后台",
    "代理H5",
    "管理后台",
    "代理域名",
    "品牌域名",
]

# 域名
DOMAIN_RE = re.compile(
    r'^(?:正常|异常|停用|过期)?\s*'
    r'([A-Za-z0-9][A-Za-z0-9.-]*\.[A-Za-z]{2,})$'
)

# 到期时间，例如：
# 2027-02-14 17:59:32(156天)
DATE_RE = re.compile(
    r'^(\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}:\d{2})'
    r'(?:\(([-+]?\d+)天\))?'
)


def classify_category(label):
    """
    将原始分类统一成图片中的分类。

    例如：
      DEMO-000-WEB                  -> WEB
      DEMO-000-WEB301海外跳转域名    -> WEB
      DEMO-000-H5301海外跳转域名     -> H5
      DEMO-000-全站APP301海外跳转域名 -> 全站APP
    """

    if not label:
        return "未分类"

    # 注意顺序，代理H5 必须在 H5 前判断
    category_keys = [
        "代理后台",
        "管理后台",
        "代理H5",
        "代理域名",
        "品牌域名",
        "全站APP",
        "WEB",
        "H5",
    ]

    for key in category_keys:
        if key in label:
            return key

    return "未分类"


def infer_category_from_header(lines, domain_index):
    """Use the first field of the line before a domain as a dynamic category."""
    for index in range(domain_index - 1, -1, -1):
        header = lines[index].strip()
        if not header:
            continue

        first_field = header.split(maxsplit=1)[0]
        if re.fullmatch(r"[A-Za-z0-9]+(?:[-_][A-Za-z0-9]+)*", first_field):
            return first_field
        break

    return "未分类"


def is_record_expired(status, expire_date, remaining_days=None, now=None):
    """Determine expiry from explicit status, absolute date, or negative days."""
    if status == "过期":
        return True

    if expire_date:
        try:
            expires_at = datetime.strptime(expire_date, "%Y-%m-%d %H:%M:%S")
            if expires_at <= (now or datetime.now()):
                return True
        except ValueError:
            pass

    return remaining_days is not None and remaining_days < 0


def parse_text(text):
    """Parse domain records from pasted or uploaded text."""
    lines = [
        line.strip()
        for line in text.splitlines()
    ]

    records = []

    for i, line in enumerate(lines):

        # 只处理正常/异常/停用/过期 开头的域名行
        if not re.match(r'^(正常|异常|停用|过期)', line):
            continue

        domain_match = DOMAIN_RE.match(line)

        if not domain_match:
            continue

        status = re.match(r'^(正常|异常|停用|过期)', line).group(1)
        domain = domain_match.group(1)

        # 取当前域名到下一个域名之间作为一个 block
        block = []

        j = i + 1

        while j < len(lines):

            next_line = lines[j]

            if (
                re.match(
                    r'^(正常|异常|停用|过期)',
                    next_line
                )
                and DOMAIN_RE.match(next_line)
            ):
                break

            block.append(next_line)

            j += 1

            # 防止异常文本无限扫描
            if len(block) > 50:
                break

        # 查找分类
        category_label = None

        for x in block:
            if re.search(
                r'-(WEB|H5|全站APP|代理后台|代理H5|管理后台|代理域名|品牌域名)',
                x
            ):
                category_label = x
                break

        category = classify_category(category_label)
        if category == "未分类":
            category = infer_category_from_header(lines, i)

        # 查找到期时间
        expire_date = ""
        expire_days = ""
        remaining_days = None

        for x in block:
            m = DATE_RE.match(x)

            if m:
                expire_date = m.group(1)

                if m.group(2) is not None:
                    remaining_days = int(m.group(2))
                    expire_days = f"{remaining_days}天"

                break

        records.append({
            "domain": domain,
            "category": category,
            "expire_date": expire_date,
            "expire_days": expire_days,
            "raw_category": category_label or "",
            "status": status,
            "is_expired": is_record_expired(
                status,
                expire_date,
                remaining_days,
            ),
        })

    return records


def parse_txt(file_path):
    """Backward-compatible file parser used by the CLI."""
    text = Path(file_path).read_text(
        encoding="utf-8-sig",
        errors="ignore"
    )
    return parse_text(text)


def create_excel(records, output_file, title="域名信息"):

    grouped = defaultdict(list)

    for item in records:
        grouped[item["category"]].append(item)

    # 主分类按照指定顺序
    categories = [
        c
        for c in CATEGORY_ORDER
        if grouped.get(c)
    ]

    # 如果碰到未知分类，也不能丢
    extra_categories = [
        c
        for c in grouped.keys()
        if c not in categories
    ]

    categories.extend(extra_categories)

    wb = Workbook()

    ws = wb.active
    ws.title = "域名信息"

    total_cols = len(categories) * 2

    # =========================
    # 第一行标题
    # =========================

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=total_cols
    )

    title_cell = ws.cell(
        row=1,
        column=1,
        value=title
    )
    # Keep a user-provided title as text rather than an Excel formula.
    title_cell.data_type = "s"

    title_cell.font = Font(
        bold=True,
        size=16
    )

    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    ws.row_dimensions[1].height = 28

    # =========================
    # 第二行表头
    # =========================

    for index, category in enumerate(categories):

        domain_col = index * 2 + 1
        expire_col = index * 2 + 2

        ws.cell(
            row=2,
            column=domain_col,
            value=category
        )

        ws.cell(
            row=2,
            column=expire_col,
            value="过期时间"
        )

    # =========================
    # 写数据
    # =========================

    max_rows = 0
    expired_font = Font(color="FFFF0000")

    for index, category in enumerate(categories):

        domain_col = index * 2 + 1
        expire_col = index * 2 + 2

        items = grouped[category]

        max_rows = max(
            max_rows,
            len(items)
        )

        for row_index, item in enumerate(
            items,
            start=3
        ):

            domain_cell = ws.cell(
                row=row_index,
                column=domain_col,
                value=item["domain"]
            )

            expire_cell = ws.cell(
                row=row_index,
                column=expire_col,
                value=item["expire_days"]
            )

            if item.get("is_expired"):
                domain_cell.font = expired_font
                expire_cell.font = expired_font

    # =========================
    # 样式
    # =========================

    thin = Side(
        style="thin",
        color="000000"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    # 标题下面所有区域加边框
    for row in ws.iter_rows(
        min_row=2,
        max_row=max_rows + 2,
        min_col=1,
        max_col=total_cols
    ):

        for cell in row:

            cell.border = border

            cell.alignment = Alignment(
                vertical="center"
            )

    # 表头加粗
    for cell in ws[2]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="left",
            vertical="center"
        )

    # =========================
    # 列宽
    # =========================

    for col in range(
        1,
        total_cols + 1
    ):

        column_letter = get_column_letter(col)

        if col % 2 == 1:
            # 域名
            ws.column_dimensions[
                column_letter
            ].width = 22

        else:
            # 过期时间
            ws.column_dimensions[
                column_letter
            ].width = 12

    # 冻结表头
    ws.freeze_panes = "A3"

    wb.save(output_file)

    return categories, grouped


def create_excel_bytes(records, title="域名信息"):
    """Create an Excel workbook in memory for an HTTP download response."""
    output = BytesIO()
    categories, grouped = create_excel(records, output, title)
    output.seek(0)
    return output, categories, grouped


def main():
    # Avoid UnicodeEncodeError in Windows terminals with a legacy code page.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser(
        description="域名TXT整理为Excel"
    )

    parser.add_argument(
        "input",
        help="TXT文件"
    )

    parser.add_argument(
        "-o",
        "--output",
        default="域名整理.xlsx",
        help="输出Excel文件"
    )

    parser.add_argument(
        "--title",
        default="域名信息",
        help="Excel顶部标题"
    )

    args = parser.parse_args()

    records = parse_txt(
        args.input
    )

    if not records:
        print("没有识别到域名")
        return

    categories, grouped = create_excel(
        records,
        args.output,
        args.title
    )

    print("=" * 50)
    print(f"总域名数量：{len(records)}")

    for category in categories:
        print(
            f"{category}: "
            f"{len(grouped[category])}"
        )

    print("=" * 50)
    print(
        f"已生成：{args.output}"
    )


if __name__ == "__main__":
    main()
